# -*- coding: utf-8 -*-
"""南京中医药大学《数据结构》课件主题与版式辅助。"""
from __future__ import annotations

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
LOGO_PATH = r"E:\游戏\符济\Rune-of-Succor\课件_数据结构_第二章线性表\assets\njucm_logo.png"

NAVY = RGBColor(0x08, 0x16, 0x2B)
NAVY2 = RGBColor(0x0C, 0x24, 0x42)
HEADER = RGBColor(0x07, 0x14, 0x28)
CARD = RGBColor(0x12, 0x34, 0x56)
CARD2 = RGBColor(0x15, 0x3D, 0x64)
ACCENT = RGBColor(0x3E, 0xC6, 0xF0)
ACCENT2 = RGBColor(0x7D, 0xD3, 0xFC)
CYAN = RGBColor(0x22, 0xD3, 0xEE)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MUTED = RGBColor(0xA9, 0xC6, 0xD9)
GOLD = RGBColor(0xE8, 0xC5, 0x47)
CODE_BG = RGBColor(0x0B, 0x10, 0x18)
OK = RGBColor(0x4A, 0xDE, 0x80)
WARN = RGBColor(0xFF, 0xB0, 0x20)
DANGER = RGBColor(0xFB, 0x71, 0x85)
SOFT = RGBColor(0x1A, 0x4A, 0x72)
LINE = RGBColor(0x2A, 0x6A, 0x96)

FONT = "微软雅黑"
FONT_CODE = "Consolas"


def C(h: str) -> RGBColor:
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def set_run_font(run, name=FONT, size=18, color=WHITE, bold=False, italic=False):
    if not isinstance(size, int):
        run.font.size = size
    else:
        run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", name)


def _fill(shape, color: RGBColor):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def _line(shape, color: RGBColor, pt=1.25):
    shape.line.color.rgb = color
    shape.line.width = Pt(pt)


def add_rect(slide, l, t, w, h, fill, name="chrome_rect", line=None, radius=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, l, t, w, h)
    _fill(sh, fill)
    if line:
        _line(sh, line[0], line[1] if len(line) > 1 else 1.25)
    else:
        sh.line.fill.background()
    if radius and hasattr(sh, "adjustments") and len(sh.adjustments) > 0:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    sh.name = name
    return sh


def add_shape(slide, kind, l, t, w, h, fill, name="chrome_shape", line=None):
    sh = slide.shapes.add_shape(kind, l, t, w, h)
    _fill(sh, fill)
    if line:
        _line(sh, line[0], line[1] if len(line) > 1 else 1.25)
    else:
        sh.line.fill.background()
    sh.name = name
    return sh


def add_textbox(slide, l, t, w, h, text, size=18, color=WHITE, bold=False,
                align="left", font=FONT, name="anim_text", anchor="top", italic=False):
    box = slide.shapes.add_textbox(l, t, w, h)
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    if anchor == "middle":
        tf.paragraphs[0].alignment = PP_ALIGN.LEFT
        box.text_frame._txBody.bodyPr.set("anchor", "ctr")
    elif anchor == "bottom":
        box.text_frame._txBody.bodyPr.set("anchor", "b")
    p = tf.paragraphs[0]
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    run = p.add_run()
    run.text = text
    set_run_font(run, font, size, color, bold, italic)
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    return box


def add_para(tf, text, size=18, color=WHITE, bold=False, align="left",
             font=FONT, space_after=8, level=0, bullet=False):
    if tf.paragraphs[0].text == "" and len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[align]
    p.level = level
    p.space_after = Pt(space_after)
    if bullet:
        pPr = p._p.get_or_add_pPr()
        bu = etree.SubElement(pPr, qn("a:buFont"))
        bu.set("typeface", "Arial")
        buChar = etree.SubElement(pPr, qn("a:buChar"))
        buChar.set("char", "•")
    run = p.add_run()
    run.text = text
    set_run_font(run, font, size, color, bold)
    return p


def add_bullets_box(slide, l, t, w, h, items, size=17, name="anim_bullets", color=WHITE):
    """items: list[str] or list[tuple(text, bold)]"""
    box = slide.shapes.add_textbox(l, t, w, h)
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        bold = False
        col = color
        if isinstance(it, tuple):
            if len(it) == 2:
                text, bold = it
            else:
                text, bold, col = it
        else:
            text = it
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        p.line_spacing = 1.15
        pPr = p._p.get_or_add_pPr()
        buFont = etree.SubElement(pPr, qn("a:buFont"))
        buFont.set("typeface", "Arial")
        buChar = etree.SubElement(pPr, qn("a:buChar"))
        buChar.set("char", "▸")
        run = p.add_run()
        run.text = " " + text
        set_run_font(run, FONT, size, col, bold)
    return box


def add_code_block(slide, l, t, w, h, code, title="C++", name="anim_code"):
    bg = add_rect(slide, l, t, w, h, CODE_BG, name="chrome_codebg", radius=0.08)
    bar = add_rect(slide, l, t, w, Inches(0.36), RGBColor(0x12, 0x1A, 0x28), name="chrome_codebar")
    add_textbox(slide, l + Inches(0.16), t + Inches(0.04), Inches(2.4), Inches(0.28),
                title, 11, ACCENT, True, name="chrome_codetitle")
    dots_x = l + w - Inches(0.9)
    for i, c in enumerate((DANGER, WARN, OK)):
        add_shape(slide, MSO_SHAPE.OVAL, dots_x + Inches(i * 0.22), t + Inches(0.1),
                  Inches(0.14), Inches(0.14), c, name="chrome_dot")
    box = slide.shapes.add_textbox(l + Inches(0.18), t + Inches(0.42), w - Inches(0.3), h - Inches(0.5))
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    lines = code.strip("\n").split("\n")
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(1)
        p.line_spacing = 1.05
        run = p.add_run()
        run.text = line if line else " "
        set_run_font(run, FONT_CODE, 12, RGBColor(0xD6, 0xE8, 0xF5), False)
    return bg


def add_tag(slide, l, t, text, fill=SOFT, fg=ACCENT2, name="chrome_tag"):
    w = Inches(max(1.15, 0.28 * len(text) + 0.35))
    h = Inches(0.32)
    add_rect(slide, l, t, w, h, fill, name=name, radius=0.5)
    add_textbox(slide, l, t, w, h, text, 11, fg, True, align="center",
                name=name + "_t", anchor="middle")
    return w


def add_card(slide, l, t, w, h, name="anim_card"):
    return add_rect(slide, l, t, w, h, CARD, name=name, radius=0.06)


def paint_background(slide):
    add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, NAVY, name="chrome_bg")
    # subtle top glow bar
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.06), ACCENT, name="chrome_topglow")
    # left accent
    add_rect(slide, 0, 0, Inches(0.07), SLIDE_H, ACCENT, name="chrome_left")


def add_logo(slide, l, t, h):
    # keep aspect 898/591 ≈ 1.52
    w = int(h * 898 / 591)
    pic = slide.shapes.add_picture(LOGO_PATH, l, t, width=w, height=h)
    pic.name = "chrome_logo"
    return pic


def add_header_footer(slide, section, page, total, kind="知识点"):
    """标准内容页页眉页脚 + 校徽。"""
    paint_background(slide)
    add_rect(slide, 0, Inches(0.06), SLIDE_W, Inches(0.78), HEADER, name="chrome_header")
    add_rect(slide, Inches(0.07), Inches(0.82), SLIDE_W - Inches(0.07), Inches(0.015),
             ACCENT, name="chrome_hline")
    add_textbox(slide, Inches(0.28), Inches(0.12), Inches(7.6), Inches(0.32),
                "南京中医药大学  ·  数据结构", 12, MUTED, False, name="chrome_school")
    add_textbox(slide, Inches(0.28), Inches(0.40), Inches(8.2), Inches(0.36),
                f"第二章 线性表  |  {section}", 16, WHITE, True, name="chrome_sect")
    add_tag(slide, Inches(9.05), Inches(0.28), kind, SOFT, ACCENT2, "chrome_kind")
    add_logo(slide, Inches(11.55), Inches(0.10), Inches(0.70))

    add_rect(slide, 0, Inches(7.18), SLIDE_W, Inches(0.32), HEADER, name="chrome_footer")
    add_textbox(slide, Inches(0.28), Inches(7.18), Inches(8.5), Inches(0.32),
                "本科课件  ·  可放映演示  ·  含课程思政与前沿拓展", 10, MUTED, name="chrome_ft1")
    add_textbox(slide, Inches(10.6), Inches(7.18), Inches(2.4), Inches(0.32),
                f"{page:02d}  /  {total:02d}", 11, ACCENT2, True, align="right", name="chrome_pg")


def add_notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def node_box(slide, l, t, w, h, text, fill=CARD2, name="anim_node", sub=None):
    sh = add_rect(slide, l, t, w, h, fill, name=name, line=(ACCENT, 1.25), radius=0.12)
    add_textbox(slide, l, t + (Inches(0.06) if sub else Inches(0.08)), w,
                Inches(0.36) if sub else h - Inches(0.1),
                text, 14 if sub else 16, WHITE, True, align="center", name=name + "_t")
    if sub:
        add_textbox(slide, l, t + Inches(0.34), w, Inches(0.28),
                    sub, 10, MUTED, False, align="center", name=name + "_s")
    return sh


def arrow_right(slide, l, t, w=Inches(0.42), h=Inches(0.18), name="anim_arrow"):
    return add_shape(slide, MSO_SHAPE.RIGHT_ARROW, l, t, w, h, ACCENT, name=name)
